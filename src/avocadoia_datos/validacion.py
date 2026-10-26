"""Planilla de validación para el experto y su incorporación (US2, research D12)."""

import os
from datetime import date, datetime

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from avocadoia_datos.clases import CLASES
from avocadoia_datos.config import ErrorPrecondicion
from avocadoia_datos.es_csv import agregar_csv

COLUMNAS_PLANILLA = ("id_contenido", "clase_propuesta", "fuente", "ruta", "decision", "clase_nueva", "objeto",
                     "toma", "calidad", "comentario", "validador", "fecha")
COLUMNAS_DECISIONES = ("id_contenido", "decision", "clase_nueva", "objeto", "toma", "calidad", "validador",
                       "fecha", "comentario", "planilla")
LISTAS = {
    "decision": ("aceptada", "rechazada", "reasignada", "mixto"),
    "clase_nueva": CLASES,
    "objeto": ("fruto_completo", "pedunculo_primer_plano", "varios_frutos", "no_es_fruto"),
    "toma": ("campo", "laboratorio", "sin_fondo"),
    "calidad": ("buena", "aceptable", "mala"),
}
_EDITABLES = ("decision", "clase_nueva", "objeto", "toma", "calidad", "comentario", "validador", "fecha")


def generar_planillas(manifiesto, raiz_imagenes, salida, clase=None, lote=250, incluir_reserva=False):
    """Escribe datos/validacion/planilla_<clase>_<NN>.xlsx y devuelve las rutas creadas."""
    if not any(f.get("en_muestra") == "si" for f in manifiesto):
        raise ErrorPrecondicion("No hay muestra: ejecute primero `particionar`.")
    carpeta = os.path.join(salida, "validacion")
    os.makedirs(carpeta, exist_ok=True)
    creadas = []
    for c in ([clase] if clase else CLASES):
        filas = [f for f in manifiesto
                 if f["canonica"] == "si" and f["clase_efectiva"] == c and f["validacion"] == "pendiente"
                 and (f.get("en_muestra") == "si" or (incluir_reserva and f["estado"] in ("elegible", "provisional")))]
        filas.sort(key=lambda f: (f.get("en_muestra") != "si", f["fuente"], f["ruta"]))
        for n, inicio in enumerate(range(0, len(filas), max(1, lote)), start=1):
            ruta = os.path.join(carpeta, f"planilla_{c}_{n:02d}.xlsx")
            _escribir_planilla(ruta, filas[inicio:inicio + lote], raiz_imagenes)
            creadas.append(ruta)
    return creadas


def _escribir_planilla(ruta, filas, raiz_imagenes):
    libro = Workbook()
    hoja = libro.active
    hoja.title = "validacion"
    hoja.append(COLUMNAS_PLANILLA)
    for celda in hoja[1]:
        celda.font = Font(bold=True)
    editable = PatternFill("solid", fgColor="FFF2CC")
    carpeta = os.path.dirname(os.path.abspath(ruta))
    for i, f in enumerate(filas, start=2):
        hoja.append([f["id_contenido"], f["clase_efectiva"], f["fuente"], f["ruta"]] + [""] * 8)
        celda = hoja.cell(row=i, column=COLUMNAS_PLANILLA.index("ruta") + 1)
        destino = os.path.relpath(os.path.abspath(os.path.join(raiz_imagenes, f["ruta"])), carpeta)
        celda.hyperlink = destino.replace(os.sep, "/")
        celda.style = "Hyperlink"
        for col in _EDITABLES:
            hoja.cell(row=i, column=COLUMNAS_PLANILLA.index(col) + 1).fill = editable
    ultima = max(2, len(filas) + 1)
    for col, valores in LISTAS.items():
        letra = hoja.cell(row=1, column=COLUMNAS_PLANILLA.index(col) + 1).column_letter
        dv = DataValidation(type="list", formula1='"' + ",".join(valores) + '"', allow_blank=True)
        dv.add(f"{letra}2:{letra}{ultima}")
        hoja.add_data_validation(dv)
    for col, ancho in {"A": 20, "B": 14, "C": 22, "D": 60, "E": 14, "F": 14, "G": 24, "H": 14, "I": 12,
                       "J": 40, "K": 18, "L": 12}.items():
        hoja.column_dimensions[col].width = ancho
    hoja.freeze_panes = "A2"
    libro.save(ruta)


def _texto(valor):
    if valor is None:
        return ""
    if isinstance(valor, (datetime, date)):
        return valor.strftime("%Y-%m-%d")
    return str(valor).strip()


def leer_planilla(ruta, ids_validos):
    """Devuelve (decisiones_validas, errores[(fila, mensaje)], ignoradas)."""
    hoja = load_workbook(ruta, data_only=True).active
    encabezado = [_texto(c.value) for c in hoja[1]]
    faltan = [c for c in COLUMNAS_PLANILLA if c not in encabezado]
    if faltan:
        return [], [(1, f"faltan columnas {faltan}")], 0
    idx = {c: encabezado.index(c) for c in COLUMNAS_PLANILLA}
    validas, errores, ignoradas = [], [], 0
    for n, fila in enumerate(hoja.iter_rows(min_row=2, values_only=True), start=2):
        v = {c: _texto(fila[idx[c]]) if idx[c] < len(fila) else "" for c in COLUMNAS_PLANILLA}
        if not v["decision"]:
            ignoradas += 1
            continue
        error = _validar_fila(v, ids_validos)
        if error:
            errores.append((n, error))
            continue
        v["planilla"] = os.path.basename(ruta)
        validas.append(v)
    return validas, errores, ignoradas


def _validar_fila(v, ids_validos):
    if v["id_contenido"] not in ids_validos:
        return f"id_contenido desconocido '{v['id_contenido']}'"
    for col, valores in LISTAS.items():
        if v[col] and v[col] not in valores:
            return f"{col}='{v[col]}' no está en {list(valores)}"
    if v["decision"] == "reasignada" and not v["clase_nueva"]:
        return "reasignada requiere clase_nueva"
    if not v["validador"]:
        return "falta validador"
    try:
        datetime.strptime(v["fecha"], "%Y-%m-%d")
    except ValueError:
        return f"fecha '{v['fecha']}' no tiene formato AAAA-MM-DD"
    return None


def incorporar(rutas, manifiesto, salida):
    ids = {f["id_contenido"] for f in manifiesto}
    resumen = {"incorporadas": 0, "ignoradas": 0, "invalidas": []}
    for ruta in rutas:
        validas, errores, ignoradas = leer_planilla(ruta, ids)
        agregar_csv(os.path.join(salida, "validacion", "decisiones.csv"), COLUMNAS_DECISIONES, validas)
        resumen["incorporadas"] += len(validas)
        resumen["ignoradas"] += ignoradas
        resumen["invalidas"] += [(os.path.basename(ruta), n, msg) for n, msg in errores]
    return resumen
