"""Planillas de revisión de los autores y su incorporación (spec 002, US1, research E4 y E11)."""

import io
import os
import random
from collections import defaultdict

from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image as ImagenXlsx
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from PIL import Image

from avocadoia_datos.clases import CLASES
from avocadoia_datos.config import ErrorPrecondicion
from avocadoia_datos.es_csv import agregar_csv, leer_csv
from avocadoia_datos.muestra import clave_estable
from avocadoia_datos.preprocesamiento import a_rgb
from avocadoia_datos.validacion import _texto

COLUMNAS_PLANILLA = ("id_contenido", "miniatura", "fuente", "grupo", "clase_asignada", "clase_sugerida", "puntaje",
                     "motivo_marca", "decision", "clase_nueva", "alcance", "atajo", "comentario", "revisor", "fecha",
                     "ruta")
COLUMNAS_DECISIONES = ("id_contenido", "decision", "clase_nueva", "alcance", "atajo", "comentario", "revisor",
                       "fecha", "planilla", "tipo")
LISTAS = {
    "decision": ("mantener", "reasignar", "excluir", "no_seguro"),
    "clase_nueva": CLASES,
    "alcance": ("grupo", "imagen"),
    "atajo": ("ninguno", "cartel", "fondo", "mano", "otro"),
}
TIPOS = ("sospecha", "sistematica", "auditoria")
_EDITABLES = ("decision", "clase_nueva", "alcance", "atajo", "comentario", "revisor", "fecha")
_MINIATURA = 110


def ruta_decisiones(salida):
    return os.path.join(salida, "curacion", "revision_etiquetas.csv")


def _decididas(salida):
    return {f["id_contenido"] for f in leer_csv(ruta_decisiones(salida))}


def generar_planillas(manifiesto, sospechas, raiz_imagenes, salida, tipo="sospecha", clase=None, lote=150):
    """Escribe datos/revision/revision_<tipo>_<clase>_<NN>.xlsx y devuelve las rutas creadas.

    - sospecha: las marcadas a_revisar, de mayor a menor puntaje, sin las ya decididas.
    - sistematica: todo el conjunto final de la clase, por fuente y ruta.
    """
    if tipo not in ("sospecha", "sistematica"):
        raise ValueError(tipo)
    if not any(f.get("en_muestra") == "si" for f in manifiesto):
        raise ErrorPrecondicion("No hay muestra: ejecute primero `todo`.")
    canonicas = {f["id_contenido"]: f for f in manifiesto if f["canonica"] == "si"}
    info = {s["id_contenido"]: s for s in sospechas}
    decididas = _decididas(salida)
    creadas = []
    for c in ([clase] if clase else CLASES):
        if tipo == "sospecha":
            filas = [canonicas[s["id_contenido"]] for s in sospechas
                     if s["a_revisar"] == "si" and s["clase_asignada"] == c and s["id_contenido"] in canonicas
                     and s["id_contenido"] not in decididas]
        else:
            filas = sorted((f for f in canonicas.values() if f["clase_efectiva"] == c and f.get("en_muestra") == "si"),
                           key=lambda f: (f["fuente"], f["ruta"]))
        # Nunca se sobrescribe una planilla existente (puede tener decisiones): se sigue la numeración.
        primero = _siguiente_numero(os.path.join(salida, "revision"), f"revision_{tipo}_{c}_")
        for n, inicio in enumerate(range(0, len(filas), max(1, lote)), start=primero):
            ruta = os.path.join(salida, "revision", f"revision_{tipo}_{c}_{n:02d}.xlsx")
            _escribir(ruta, filas[inicio:inicio + lote], info, raiz_imagenes)
            creadas.append(ruta)
    return creadas


def _siguiente_numero(carpeta, prefijo):
    usados = [int(n[len(prefijo):len(prefijo) + 2]) for n in (os.listdir(carpeta) if os.path.isdir(carpeta) else [])
              if n.startswith(prefijo) and n[len(prefijo):len(prefijo) + 2].isdigit()]
    return max(usados, default=0) + 1


def generar_auditoria(particion, manifiesto, raiz_imagenes, salida, n=100, semilla=42):
    """Muestra estratificada por clase del conjunto final → datos/revision/auditoria_01.xlsx."""
    finales = [f for f in particion if f["particion"] in ("train", "val", "test")]
    if not finales:
        raise ErrorPrecondicion("No hay partición: ejecute primero `todo`.")
    canonicas = {f["id_contenido"]: f for f in manifiesto if f["canonica"] == "si"}
    por_clase = defaultdict(list)
    for f in finales:
        por_clase[f["clase"]].append(f["id_contenido"])
    cuotas = _mayor_resto(n, {c: len(v) for c, v in por_clase.items()})
    elegidas = []
    for c in CLASES:
        orden = sorted(por_clase.get(c, []), key=lambda i: (clave_estable(semilla, "auditoria", i), i))
        elegidas += orden[:cuotas.get(c, 0)]
    # Orden aleatorio estable, para que el revisor no vea las clases en bloque.
    random.Random(semilla).shuffle(elegidas)
    ruta = os.path.join(salida, "revision", f"auditoria_{_siguiente_numero(os.path.join(salida, 'revision'), 'auditoria_'):02d}.xlsx")
    _escribir(ruta, [canonicas[i] for i in elegidas], {}, raiz_imagenes)
    return ruta


def _mayor_resto(total, pesos):
    suma = sum(pesos.values())
    exactas = {k: total * v / suma for k, v in pesos.items()}
    cuotas = {k: int(v) for k, v in exactas.items()}
    for k in sorted(exactas, key=lambda k: (-(exactas[k] - cuotas[k]), k))[:total - sum(cuotas.values())]:
        cuotas[k] += 1
    return cuotas


def _miniatura(ruta):
    with Image.open(ruta) as im:
        im = a_rgb(im)
        im.thumbnail((_MINIATURA, _MINIATURA))
        buf = io.BytesIO()
        im.save(buf, "PNG")
    buf.seek(0)
    return ImagenXlsx(buf)


def _escribir(ruta, filas, info, raiz_imagenes):
    libro = Workbook()
    hoja = libro.active
    hoja.title = "revision"
    hoja.append(COLUMNAS_PLANILLA)
    for celda in hoja[1]:
        celda.font = Font(bold=True)
    editable = PatternFill("solid", fgColor="FFF2CC")
    carpeta = os.path.dirname(os.path.abspath(ruta))
    col = {c: COLUMNAS_PLANILLA.index(c) + 1 for c in COLUMNAS_PLANILLA}
    for i, f in enumerate(filas, start=2):
        s = info.get(f["id_contenido"], {})
        hoja.append([f["id_contenido"], "", f["fuente"], f["grupo"], f["clase_efectiva"], s.get("clase_sugerida", ""),
                     float(s["puntaje"]) if s.get("puntaje") else "", s.get("motivo_marca", ""),
                     "", "", "grupo", "", "", "", "", f["ruta"]])
        hoja.row_dimensions[i].height = _MINIATURA * 0.78
        original = os.path.join(raiz_imagenes, f["ruta"])
        if os.path.exists(original):
            hoja.add_image(_miniatura(original), hoja.cell(row=i, column=col["miniatura"]).coordinate)
        enlace = hoja.cell(row=i, column=col["ruta"])
        enlace.hyperlink = os.path.relpath(os.path.abspath(original), carpeta).replace(os.sep, "/")
        enlace.style = "Hyperlink"
        for c in _EDITABLES:
            hoja.cell(row=i, column=col[c]).fill = editable
        for c in COLUMNAS_PLANILLA:
            hoja.cell(row=i, column=col[c]).alignment = Alignment(vertical="center", wrap_text=c == "comentario")
    ultima = max(2, len(filas) + 1)
    for c, valores in LISTAS.items():
        letra = hoja.cell(row=1, column=col[c]).column_letter
        dv = DataValidation(type="list", formula1='"' + ",".join(valores) + '"', allow_blank=True)
        dv.add(f"{letra}2:{letra}{ultima}")
        hoja.add_data_validation(dv)
    anchos = {"id_contenido": 19, "miniatura": 16, "fuente": 20, "grupo": 18, "clase_asignada": 13,
              "clase_sugerida": 13, "puntaje": 8, "motivo_marca": 16, "decision": 12, "clase_nueva": 13,
              "alcance": 9, "atajo": 10, "comentario": 30, "revisor": 14, "fecha": 11, "ruta": 50}
    for c, ancho in anchos.items():
        hoja.column_dimensions[hoja.cell(row=1, column=col[c]).column_letter].width = ancho
    hoja.freeze_panes = "C2"
    os.makedirs(os.path.dirname(ruta), exist_ok=True)
    libro.save(ruta)


def _tipo_de(nombre):
    if nombre.startswith("auditoria"):
        return "auditoria"
    partes = nombre.split("_")
    return partes[1] if len(partes) > 1 and partes[1] in TIPOS else "sospecha"


def validar_fila(v, ids_validos):
    if v["id_contenido"] not in ids_validos:
        return f"id_contenido desconocido '{v['id_contenido']}'"
    for c, valores in LISTAS.items():
        if v[c] and v[c] not in valores:
            return f"{c}='{v[c]}' no está en {list(valores)}"
    if v["decision"] == "reasignar" and not v["clase_nueva"]:
        return "reasignar requiere clase_nueva"
    if not v["revisor"]:
        return "falta revisor"
    if len(v["fecha"]) != 10 or v["fecha"][4] != "-" or v["fecha"][7] != "-" or not v["fecha"].replace("-", "").isdigit():
        return f"fecha '{v['fecha']}' no tiene formato AAAA-MM-DD"
    return None


def leer_planilla(ruta, ids_validos):
    """Devuelve (decisiones_validas, errores[(fila, mensaje)], ignoradas)."""
    hoja = load_workbook(ruta, data_only=True).active
    encabezado = [_texto(c.value) for c in hoja[1]]
    faltan = [c for c in COLUMNAS_PLANILLA if c not in encabezado]
    if faltan:
        return [], [(1, f"faltan columnas {faltan}")], 0
    idx = {c: encabezado.index(c) for c in COLUMNAS_PLANILLA}
    tipo = _tipo_de(os.path.basename(ruta))
    validas, errores, ignoradas = [], [], 0
    for n, fila in enumerate(hoja.iter_rows(min_row=2, values_only=True), start=2):
        v = {c: _texto(fila[idx[c]]) if idx[c] < len(fila) else "" for c in COLUMNAS_PLANILLA}
        if not v["decision"]:
            ignoradas += 1
            continue
        v["alcance"] = v["alcance"] or "grupo"
        error = validar_fila(v, ids_validos)
        if error:
            errores.append((n, error))
            continue
        validas.append({**{c: v[c] for c in COLUMNAS_DECISIONES if c in v},
                        "planilla": os.path.basename(ruta), "tipo": tipo})
    return validas, errores, ignoradas


def incorporar(rutas, manifiesto, salida):
    ids = {f["id_contenido"] for f in manifiesto}
    resumen = {"incorporadas": 0, "ignoradas": 0, "invalidas": []}
    for ruta in rutas:
        validas, errores, ignoradas = leer_planilla(ruta, ids)
        agregar_csv(ruta_decisiones(salida), COLUMNAS_DECISIONES, validas)
        resumen["incorporadas"] += len(validas)
        resumen["ignoradas"] += ignoradas
        resumen["invalidas"] += [(os.path.basename(ruta), n, msg) for n, msg in errores]
    return resumen


def decisiones_vigentes(filas_decision, grupo_de):
    """{id_contenido: decisión} con la regla de alcance (data-model, Decisión de revisión).

    Gana la decisión más reciente (fecha y luego orden en el archivo). Una decisión de alcance
    `grupo` alcanza a todos los miembros del grupo de la imagen decidida.
    """
    miembros = defaultdict(set)
    for i, g in grupo_de.items():
        miembros[g].add(i)
    vigentes = {}
    for orden, d in enumerate(filas_decision):
        clave = (d.get("fecha", ""), orden)
        alcanzadas = miembros.get(grupo_de.get(d["id_contenido"]), set()) if d.get("alcance", "grupo") == "grupo" else set()
        for i in alcanzadas | {d["id_contenido"]}:
            if i not in vigentes or clave >= vigentes[i][0]:
                vigentes[i] = (clave, d)
    return {i: v[1] for i, v in vigentes.items()}


def grupos_partidos(filas, grupos_revisados):
    """Entre los grupos con alguna decisión de alcance `imagen`, los que quedaron con clases o
    estados de revisión distintos."""
    por_grupo = defaultdict(set)
    for f in filas:
        if f["canonica"] == "si" and f["grupo"] in grupos_revisados:
            por_grupo[f["grupo"]].add((f["clase_efectiva"], f["estado"] == "excluida_revision"))
    return sorted(g for g, v in por_grupo.items() if len(v) > 1)
