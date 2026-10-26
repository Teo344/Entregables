import os

import pytest
from openpyxl import load_workbook

from avocadoia_datos import inventario, particion, validacion
from avocadoia_datos.config import ErrorPrecondicion
from avocadoia_datos.es_csv import leer_csv


def _preparar(arbol):
    cfg = arbol.config()
    m = inventario.ejecutar(cfg, arbol.modelo, arbol.salida, arbol.confirmados())[0]
    particion.ejecutar(m, cfg, arbol.salida)
    return cfg, inventario.leer_manifiesto(arbol.salida)


def test_planilla_requiere_muestra(arbol):
    with pytest.raises(ErrorPrecondicion):
        validacion.generar_planillas([], arbol.modelo, arbol.salida)


def test_planilla_lotes_enlaces_y_listas(arbol_completo):
    _, m = _preparar(arbol_completo)
    rutas = validacion.generar_planillas(m, arbol_completo.modelo, arbol_completo.salida, "pedunculo", lote=15)
    assert [os.path.basename(r) for r in rutas] == ["planilla_pedunculo_01.xlsx", "planilla_pedunculo_02.xlsx",
                                                    "planilla_pedunculo_03.xlsx"]
    hoja = load_workbook(rutas[0]).active
    assert [c.value for c in hoja[1]] == list(validacion.COLUMNAS_PLANILLA)
    assert hoja.max_row == 16
    enlace = hoja.cell(row=2, column=4).hyperlink.target
    assert os.path.exists(os.path.join(os.path.dirname(rutas[0]), enlace))
    formulas = {dv.formula1 for dv in hoja.data_validations.dataValidation}
    assert '"aceptada,rechazada,reasignada,mixto"' in formulas
    assert '"fruto_completo,pedunculo_primer_plano,varios_frutos,no_es_fruto"' in formulas


def test_incorporar_validas_invalidas_y_vacias(arbol_completo):
    _, m = _preparar(arbol_completo)
    ruta = validacion.generar_planillas(m, arbol_completo.modelo, arbol_completo.salida, "pedunculo", lote=10)[0]
    libro = load_workbook(ruta)
    hoja = libro.active
    col = {c.value: i + 1 for i, c in enumerate(hoja[1])}

    def poner(fila, **valores):
        for k, v in valores.items():
            hoja.cell(row=fila, column=col[k]).value = v

    poner(2, decision="aceptada", validador="Ing. Sigcha", fecha="2026-10-25", objeto="pedunculo_primer_plano")
    poner(3, decision="rechazada", validador="Ing. Sigcha", fecha="2026-10-25")
    poner(4, decision="reasignada", clase_nueva="antracnosis", validador="Ing. Sigcha", fecha="2026-10-25")
    poner(5, decision="quizas", validador="Ing. Sigcha", fecha="2026-10-25")
    poner(6, decision="reasignada", validador="Ing. Sigcha", fecha="2026-10-25")
    poner(7, decision="aceptada", fecha="2026-10-25")
    poner(8, decision="aceptada", validador="x", fecha="25/10/2026")
    libro.save(ruta)

    r = validacion.incorporar([ruta], m, arbol_completo.salida)
    assert r["incorporadas"] == 3
    assert r["ignoradas"] == 3
    assert [n for _, n, _ in r["invalidas"]] == [5, 6, 7, 8]
    decisiones = leer_csv(os.path.join(arbol_completo.salida, "validacion", "decisiones.csv"))
    assert [d["decision"] for d in decisiones] == ["aceptada", "rechazada", "reasignada"]
    assert decisiones[0]["planilla"] == "planilla_pedunculo_01.xlsx"

    # Al reinventariar, la reasignada cambia de clase y la rechazada sale de la muestra al particionar.
    cfg = arbol_completo.config()
    m2 = inventario.ejecutar(cfg, arbol_completo.modelo, arbol_completo.salida, arbol_completo.confirmados())[0]
    por_id = {f["id_contenido"]: f for f in m2}
    assert por_id[decisiones[1]["id_contenido"]]["estado"] == "rechazada"
    assert por_id[decisiones[2]["id_contenido"]]["clase_efectiva"] == "antracnosis"
    filas, _ = particion.ejecutar(m2, cfg, arbol_completo.salida)
    assert decisiones[1]["id_contenido"] not in {f["id_contenido"] for f in filas}
