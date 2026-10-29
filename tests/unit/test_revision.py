"""Planillas de revisión de los autores e incorporación (spec 002, T014)."""

import os

from openpyxl import load_workbook

from avocadoia_datos import inventario
from avocadoia_datos.es_csv import leer_csv
from avocadoia_datos.etl import revision


def _manifiesto(arbol):
    for fruto in range(3):
        for dia in (1, 2):
            arbol.agregar(f"AGUACATESANOS/hass_T10_d{dia:02d}_{fruto:03d}_a_1_x{fruto}{dia}.jpg")
    arbol.agregar("ANTRACNOSIS/rgb_antracnosis_0001.jpg")
    filas, _ = inventario.ejecutar(arbol.config(), arbol.modelo, arbol.salida, arbol.confirmados())
    for f in filas:  # simula una muestra ya seleccionada
        f["en_muestra"] = "si"
    return filas


def _sospechas(filas):
    return [{"id_contenido": f["id_contenido"], "clase_asignada": f["clase_efectiva"], "clase_sugerida": "antracnosis",
             "puntaje": "0.9000", "motivo_marca": "conflicto", "a_revisar": "si"} for f in filas]


def _completar(ruta, valores_por_fila):
    libro = load_workbook(ruta)
    hoja = libro.active
    col = {c.value: c.column for c in hoja[1]}
    for fila, valores in valores_por_fila.items():
        for c, v in valores.items():
            hoja.cell(row=fila, column=col[c], value=v)
    libro.save(ruta)


def test_planilla_tiene_las_columnas_y_miniaturas(arbol):
    filas = _manifiesto(arbol)
    rutas = revision.generar_planillas(filas, _sospechas(filas), arbol.modelo, arbol.salida, clase="sano")
    assert [os.path.basename(r) for r in rutas] == ["revision_sospecha_sano_01.xlsx"]
    hoja = load_workbook(rutas[0]).active
    assert [c.value for c in hoja[1]] == list(revision.COLUMNAS_PLANILLA)
    assert hoja.max_row == 7  # 6 fotos hass
    assert len(hoja._images) == 6
    assert hoja.cell(row=2, column=revision.COLUMNAS_PLANILLA.index("alcance") + 1).value == "grupo"


def test_incorporar_rechaza_filas_invalidas_y_acepta_las_validas(arbol):
    filas = _manifiesto(arbol)
    ruta = revision.generar_planillas(filas, _sospechas(filas), arbol.modelo, arbol.salida, clase="sano")[0]
    base = {"revisor": "M. Criollo", "fecha": "2026-10-27"}
    _completar(ruta, {
        2: {"decision": "quizas", **base},
        3: {"decision": "reasignar", **base},
        4: {"decision": "excluir", "revisor": "M. Criollo", "fecha": "27/10/2026"},
        5: {"decision": "excluir", "fecha": "2026-10-27"},
        6: {"decision": "excluir", "comentario": "sin fruto", **base},
    })
    r = revision.incorporar([ruta], filas, arbol.salida)
    mensajes = [m for _, _, m in r["invalidas"]]
    assert r["incorporadas"] == 1 and r["ignoradas"] == 1 and len(mensajes) == 4
    assert any("decision" in m for m in mensajes)
    assert any("clase_nueva" in m for m in mensajes)
    assert any("AAAA-MM-DD" in m for m in mensajes)
    assert any("revisor" in m for m in mensajes)
    guardada = leer_csv(revision.ruta_decisiones(arbol.salida))[0]
    assert guardada["tipo"] == "sospecha" and guardada["alcance"] == "grupo"


def test_gana_la_mas_reciente_y_el_grupo_se_propaga():
    grupo_de = {"a": "g1", "b": "g1", "c": "g2"}
    decisiones = [
        {"id_contenido": "a", "decision": "excluir", "alcance": "grupo", "fecha": "2026-10-27"},
        {"id_contenido": "b", "decision": "mantener", "alcance": "imagen", "fecha": "2026-10-28"},
        {"id_contenido": "c", "decision": "excluir", "alcance": "imagen", "fecha": "2026-10-28"},
        {"id_contenido": "c", "decision": "mantener", "alcance": "imagen", "fecha": "2026-10-27"},
    ]
    vigentes = revision.decisiones_vigentes(decisiones, grupo_de)
    assert vigentes["a"]["decision"] == "excluir"
    assert vigentes["b"]["decision"] == "mantener"  # su decisión propia es más reciente
    assert vigentes["c"]["decision"] == "excluir"   # gana la fecha, no el orden


def test_alcance_imagen_que_parte_un_grupo_genera_advertencia(arbol):
    filas = _manifiesto(arbol)
    hass = sorted((f for f in filas if f["fuente"] == "hass"), key=lambda f: f["ruta"])
    ruta = revision.generar_planillas(filas, _sospechas(filas), arbol.modelo, arbol.salida, clase="sano")[0]
    hoja = load_workbook(ruta).active
    fila_de = {hoja.cell(row=r, column=1).value: r for r in range(2, hoja.max_row + 1)}
    _completar(ruta, {fila_de[hass[0]["id_contenido"]]: {"decision": "excluir", "alcance": "imagen",
                                                          "revisor": "J. Yépez", "fecha": "2026-10-27"}})
    revision.incorporar([ruta], filas, arbol.salida)
    _, meta = inventario.ejecutar(arbol.config(), arbol.modelo, arbol.salida, arbol.confirmados())
    assert meta["grupos_partidos"] == [hass[0]["grupo"]]


def test_no_sobrescribe_planillas_existentes(arbol):
    filas = _manifiesto(arbol)
    primera = revision.generar_planillas(filas, _sospechas(filas), arbol.modelo, arbol.salida, clase="sano")
    segunda = revision.generar_planillas(filas, _sospechas(filas), arbol.modelo, arbol.salida, clase="sano")
    assert [os.path.basename(r) for r in primera + segunda] == ["revision_sospecha_sano_01.xlsx",
                                                                "revision_sospecha_sano_02.xlsx"]
