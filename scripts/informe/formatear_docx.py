"""Formato del Word generado por pandoc: Calibri, tablas con bordes y encabezado sombreado.

Uso: python scripts/informe/formatear_docx.py ENTRADA.docx SALIDA.docx
(migrado del informe de la fase 1; spec 002, T044)
"""

import sys

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


def formatear(entrada, salida):
    doc = Document(entrada)
    for s in doc.styles:
        try:
            if s.type == 1 and s.font is not None:
                s.font.name = "Calibri"
                if s.name in ("Normal", "Body Text", "First Paragraph", "Compact"):
                    s.font.size = Pt(11)
        except Exception:
            pass
    for t in doc.tables:
        bordes = OxmlElement("w:tblBorders")
        for e in ("top", "left", "bottom", "right", "insideH", "insideV"):
            b = OxmlElement(f"w:{e}")
            b.set(qn("w:val"), "single")
            b.set(qn("w:sz"), "4")
            b.set(qn("w:color"), "999999")
            bordes.append(b)
        t._tbl.tblPr.append(bordes)
        for c in t.rows[0].cells:
            sombra = OxmlElement("w:shd")
            sombra.set(qn("w:val"), "clear")
            sombra.set(qn("w:fill"), "E8F0E0")
            c._tc.get_or_add_tcPr().append(sombra)
            for p in c.paragraphs:
                for r in p.runs:
                    r.bold = True
        for fila in t.rows:
            for c in fila.cells:
                for p in c.paragraphs:
                    for r in p.runs:
                        r.font.size = Pt(9.5)
    # Las figuras ocupan como máximo el ancho útil de la página.
    for forma in doc.inline_shapes:
        if forma.width > Inches(6.3):
            proporcion = forma.height / forma.width
            forma.width = Inches(6.3)
            forma.height = int(Inches(6.3) * proporcion)
    doc.save(salida)
    print("tablas:", len(doc.tables), "| figuras:", len(doc.inline_shapes))


if __name__ == "__main__":
    formatear(sys.argv[1], sys.argv[2])
